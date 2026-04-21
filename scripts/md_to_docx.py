"""Convert docs/project_evaluation_framework.md to .docx."""
from pathlib import Path
import re

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

SRC = Path("docs/project_evaluation_framework.md")
DST = Path("docs/project_evaluation_framework.docx")


def add_runs(paragraph, text: str) -> None:
    """Render inline bold/italic markdown into runs."""
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        else:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    """Read a pipe-delimited markdown table starting at `start`. Returns (rows, next_index)."""
    rows = []
    i = start
    while i < len(lines) and lines[i].lstrip().startswith("|"):
        row = lines[i].strip()
        cells = [c.strip() for c in row.strip("|").split("|")]
        # Skip the header-separator row (|---|---|)
        if not all(re.fullmatch(r":?-+:?", c) for c in cells):
            rows.append(cells)
        i += 1
    return rows, i


def main() -> None:
    text = SRC.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Horizontal rule
        if stripped == "---":
            doc.add_paragraph().add_run().add_break()
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            doc.add_heading(stripped[2:].strip(), level=0)
            i += 1
            continue
        if stripped.startswith("## "):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue
        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        # Table
        if stripped.startswith("|"):
            rows, i = parse_table(lines, i)
            if rows:
                table = doc.add_table(rows=len(rows), cols=len(rows[0]))
                table.style = "Light Grid Accent 1"
                for r_idx, row in enumerate(rows):
                    for c_idx, cell in enumerate(row):
                        tcell = table.cell(r_idx, c_idx)
                        tcell.text = ""
                        p = tcell.paragraphs[0]
                        add_runs(p, cell)
                        if r_idx == 0:
                            for run in p.runs:
                                run.bold = True
                doc.add_paragraph()
            continue

        # Bullet list
        if stripped.startswith("- [ ] "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, "☐ " + stripped[6:])
            i += 1
            continue
        if stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, stripped[2:])
            i += 1
            continue

        # Numbered list
        m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_runs(p, m.group(2))
            i += 1
            continue

        # Italic-only line (e.g., "*Version-by-version...*")
        if stripped.startswith("*") and stripped.endswith("*") and not stripped.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(stripped.strip("*"))
            run.italic = True
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Plain paragraph
        p = doc.add_paragraph()
        add_runs(p, stripped)
        i += 1

    doc.save(DST)
    print(f"Wrote {DST} ({DST.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
